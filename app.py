import os
import sys
import time
import socket
from datetime import datetime
from io import BytesIO
from werkzeug.utils import secure_filename
from flask import Flask, render_template, request, jsonify, redirect, url_for, session, send_file, send_from_directory
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Windows konsolunda Türkçe ve emojilerin hatasız yazılması için UTF-8 yapılandırması
if sys.platform == 'win32':
    try:
        if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
            sys.stdout.reconfigure(encoding='utf-8', errors='replace')
        if sys.stderr and hasattr(sys.stderr, 'reconfigure'):
            sys.stderr.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

from dotenv import load_dotenv

load_dotenv()

import database

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'change-this-secret-key')

# Operasyon Yetkilendirme Bilgileri (Ortam Değişkeni veya Varsayılan)
OPERATIONS_USERNAME = os.environ.get('OPERATIONS_USERNAME', 'sistem')
OPERATIONS_PASSWORD = os.environ.get('OPERATIONS_PASSWORD', 'change-this-operations-password')
ADMIN_USERNAME = os.environ.get('ADMIN_USERNAME', 'admin')
ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', 'change-this-admin-password')

def get_operations_password():
    """Veritabanından güncel şifreyi okur, yoksa ortam değişkeni veya varsayılanı döner"""
    stored = database.get_setting('operations_password')
    if stored:
        return stored
    return OPERATIONS_PASSWORD

# Dosya Yükleme Klasörü (T1 / Tır Karnesi Fotoğrafları)
UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # Maksimum 16MB dosya

# Veritabanını başlat
database.init_db()

# Kimlik doğrulama kontrolü
def is_authenticated():
    return session.get('logged_in') is True

def get_local_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        # Bağlantı kurmadan yerel IP tespit etme
        s.connect(('8.8.8.8', 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return '127.0.0.1'

MAPS_URL = os.environ.get('MAPS_URL', '')

@app.route('/uploads/<path:filename>')
def uploaded_file(filename):
    """Yüklenen belge fotoğraflarını sunar"""
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)

# ==========================================
# SAYFA ROTALARI
# ==========================================

@app.route('/')
@app.route('/kayit')
def driver_form():
    """Herkese açık şoför / nakliyeci araç kayıt formu"""
    driver_notice = database.get_setting('driver_notice', '')
    return render_template('kayit.html', maps_url=MAPS_URL, driver_notice=driver_notice)

@app.route('/login')
def login_page():
    """Operasyon giriş sayfası"""
    if is_authenticated():
        return redirect(url_for('operations_dashboard'))
    return render_template('login.html')

@app.route('/operasyon')
def operations_dashboard():
    """Operasyon paneli (Masaüstü & Mobil uyumlu)"""
    if not is_authenticated():
        return redirect(url_for('login_page'))
    return render_template('operasyon.html')

# ==========================================
# API ROTALARI
# ==========================================

@app.route('/api/kayit', methods=['POST'])
def api_add_vehicle():
    """Şoför / Nakliyeci tarafından herkese açık araç girişi (Fotoğraf destekli)"""
    photo_file = None
    if request.is_json:
        data = request.get_json() or {}
    else:
        data = request.form.to_dict()
        photo_file = request.files.get('document_photo')
    
    tractor_plate = data.get('tractor_plate', '').strip()
    trailer_plate = data.get('trailer_plate', '').strip()
    company_name = data.get('company_name', '').strip()
    arrival_date = data.get('arrival_date', '').strip()

    # Zorunlu alan kontrolü
    if not tractor_plate or not trailer_plate or not company_name or not arrival_date:
        return jsonify({
            'success': False,
            'message': 'Lütfen zorunlu alanları doldurunuz (Çekici Plaka, Dorse Plaka, Firma Adı, Varış Tarihi).'
        }), 400

    try:
        gross_weight = float(data.get('gross_weight')) if data.get('gross_weight') else None
    except (ValueError, TypeError):
        gross_weight = None

    try:
        net_weight = float(data.get('net_weight')) if data.get('net_weight') else None
    except (ValueError, TypeError):
        net_weight = None

    try:
        kantar_weight = float(data.get('kantar_weight')) if data.get('kantar_weight') else None
    except (ValueError, TypeError):
        kantar_weight = None

    try:
        package_count = int(data.get('package_count')) if data.get('package_count') else None
    except (ValueError, TypeError):
        package_count = None

    customs_doc_no = data.get('customs_doc_no', '').strip()
    important_notes = data.get('important_notes', '').strip()

    # Taşıma Belgesi Fotoğrafı Yükleme (T1 / Tır Karnesi)
    document_photo_url = data.get('document_photo', '').strip()
    if photo_file and photo_file.filename:
        filename = secure_filename(photo_file.filename)
        ext = os.path.splitext(filename)[1].lower()
        if ext in ['.jpg', '.jpeg', '.png', '.webp', '.pdf', '.heic']:
            safe_plate = "".join(c for c in tractor_plate if c.isalnum())
            unique_name = f"doc_{int(time.time())}_{safe_plate}{ext}"
            file_path = os.path.join(app.config['UPLOAD_FOLDER'], unique_name)
            photo_file.save(file_path)
            document_photo_url = f"/uploads/{unique_name}"

    vehicle_id = database.add_vehicle(
        tractor_plate=tractor_plate,
        trailer_plate=trailer_plate,
        company_name=company_name,
        arrival_date=arrival_date,
        gross_weight=gross_weight,
        package_count=package_count,
        customs_doc_no=customs_doc_no,
        important_notes=important_notes,
        document_photo=document_photo_url,
        net_weight=net_weight,
        kantar_weight=kantar_weight
    )

    return jsonify({
        'success': True,
        'id': vehicle_id,
        'document_photo': document_photo_url,
        'message': f'{tractor_plate} plakalı araç kaydı başarıyla oluşturuldu.'
    }), 201

@app.route('/api/settings/notice', methods=['GET'])
def api_get_driver_notice():
    """Şoförler için aktif önemli not duyurusunu getir"""
    notice = database.get_setting('driver_notice', '')
    return jsonify({'success': True, 'notice': notice, 'maps_url': MAPS_URL})

@app.route('/api/settings/notice', methods=['POST'])
def api_set_driver_notice():
    """Operatör tarafından şoför önemli not duyurusunu güncelle"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401
    data = request.get_json() or {}
    notice = data.get('notice', '').strip()
    database.set_setting('driver_notice', notice)
    return jsonify({'success': True, 'message': 'Şoförler için önemli not başarıyla güncellendi.'})

@app.route('/api/login', methods=['POST'])
def api_login():
    """Operasyon ekibi girişi (sistem / dinamik operasyon şifresi)"""
    data = request.get_json() or {}
    username = data.get('username', '').strip()
    password = data.get('password', '').strip()

    valid_password = get_operations_password()
    if username == OPERATIONS_USERNAME and password == valid_password:
        session['logged_in'] = True
        session['user'] = username
        return jsonify({'success': True, 'message': 'Giriş başarılı.'})
    else:
        return jsonify({'success': False, 'message': 'Hatalı kullanıcı adı veya şifre.'}), 401

@app.route('/api/admin/change-password', methods=['POST'])
def api_admin_change_password():
    """Admin kimlik doğrulaması ile operasyon ('sistem') kullanıcısının şifresini değiştirme"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401

    data = request.get_json() or {}
    admin_user = data.get('admin_username', '').strip()
    admin_pass = data.get('admin_password', '').strip()
    new_password = data.get('new_password', '').strip()
    new_password_confirm = data.get('new_password_confirm', '').strip()

    # Yönetici kullanıcı adı ve şifre kontrolü
    if admin_user != ADMIN_USERNAME or admin_pass != ADMIN_PASSWORD:
        return jsonify({
            'success': False,
            'message': 'Yetkisiz erişim! Yönetici kullanıcı adı veya yönetici şifresi hatalı.'
        }), 403

    if not new_password:
        return jsonify({'success': False, 'message': 'Yeni şifre boş bırakılamaz.'}), 400

    if len(new_password) < 4:
        return jsonify({'success': False, 'message': 'Yeni şifre en az 4 karakter olmalıdır.'}), 400

    if new_password != new_password_confirm:
        return jsonify({'success': False, 'message': 'Girilen yeni şifreler birbiriyle uyuşmuyor.'}), 400

    # Veritabanında güncelle
    database.set_setting('operations_password', new_password)
    return jsonify({
        'success': True,
        'message': f"Operasyon ('{OPERATIONS_USERNAME}') giriş şifresi başarıyla güncellendi."
    })

@app.route('/api/logout', methods=['POST'])
def api_logout():
    session.clear()
    return jsonify({'success': True})

@app.route('/api/vehicles', methods=['GET'])
def api_get_vehicles():
    """Araç listesi (Filtrelemeli)"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401

    status = request.args.get('status')
    company = request.args.get('company')
    date_start = request.args.get('date_start')
    date_end = request.args.get('date_end')
    search = request.args.get('search')
    is_active_only = request.args.get('is_active', 'false').lower() == 'true'

    vehicles = database.get_vehicles(
        status=status,
        company=company,
        date_start=date_start,
        date_end=date_end,
        search=search,
        is_active_only=is_active_only
    )
    return jsonify({'success': True, 'vehicles': vehicles})

@app.route('/api/vehicles/<int:vehicle_id>', methods=['GET'])
def api_get_vehicle_detail(vehicle_id):
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401
        
    v = database.get_vehicle_by_id(vehicle_id)
    if not v:
        return jsonify({'error': 'Not found'}), 404
    return jsonify({'success': True, 'vehicle': v})

@app.route('/api/vehicles/<int:vehicle_id>', methods=['PUT'])
def api_update_vehicle(vehicle_id):
    """Araç bilgilerini güncelleme (Operatör tarafından)"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401

    data = request.get_json() or {}
    
    # Sayısal dönüşümler
    if 'gross_weight' in data:
        try:
            data['gross_weight'] = float(data['gross_weight']) if data['gross_weight'] != '' and data['gross_weight'] is not None else None
        except (ValueError, TypeError):
            data['gross_weight'] = None

    if 'net_weight' in data:
        try:
            data['net_weight'] = float(data['net_weight']) if data['net_weight'] != '' and data['net_weight'] is not None else None
        except (ValueError, TypeError):
            data['net_weight'] = None

    if 'kantar_weight' in data:
        try:
            data['kantar_weight'] = float(data['kantar_weight']) if data['kantar_weight'] != '' and data['kantar_weight'] is not None else None
        except (ValueError, TypeError):
            data['kantar_weight'] = None

    if 'package_count' in data:
        try:
            data['package_count'] = int(data['package_count']) if data['package_count'] != '' and data['package_count'] is not None else None
        except (ValueError, TypeError):
            data['package_count'] = None

    if 'is_registered' in data:
        try:
            data['is_registered'] = int(data['is_registered'])
            if data['is_registered'] == 1 and not data.get('registered_at'):
                data['registered_at'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            elif data['is_registered'] == 0:
                data['registered_at'] = None
        except (ValueError, TypeError):
            pass

    success = database.update_vehicle(vehicle_id, **data)
    if success:
        return jsonify({'success': True, 'message': 'Araç bilgileri güncellendi.'})
    return jsonify({'success': False, 'message': 'Güncelleme başarısız oldu.'}), 400

@app.route('/api/vehicles/<int:vehicle_id>/tescil', methods=['POST'])
def api_toggle_tescil(vehicle_id):
    """Tescil yapıldı mı durumunu aç/kapat"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401

    data = request.get_json() or {}
    force_state = data.get('state')
    
    result = database.toggle_tescil(vehicle_id, force_state)
    if result:
        return jsonify({'success': True, 'data': result})
    return jsonify({'success': False, 'message': 'İşlem başarısız.'}), 400

@app.route('/api/vehicles/<int:vehicle_id>/liman', methods=['POST'])
def api_toggle_liman(vehicle_id):
    """Liman giriş kaydı yapıldı mı durumunu aç/kapat"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401

    data = request.get_json() or {}
    force_state = data.get('state')
    
    result = database.toggle_port_entry(vehicle_id, force_state)
    if result:
        return jsonify({'success': True, 'data': result})
    return jsonify({'success': False, 'message': 'İşlem başarısız.'}), 400

@app.route('/api/vehicles/<int:vehicle_id>/bosaltma', methods=['POST'])
def api_complete_unloading(vehicle_id):
    """Boşaltma tamamlama ve onaylama (miktar opsiyonel, net ağırlık opsiyonel, kantar miktarı opsiyonel)"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401

    data = request.get_json() or {}
    unloaded_amount = data.get('unloaded_amount', '').strip()
    if not unloaded_amount:
        unloaded_amount = 'Onaylandı'
    unloaded_note = data.get('unloaded_note', '').strip()
    dock_number = data.get('dock_number', '').strip() if data.get('dock_number') else None

    net_weight_val = None
    if 'net_weight' in data and data['net_weight'] is not None and str(data['net_weight']).strip() != '':
        try:
            net_weight_val = float(data['net_weight'])
        except (ValueError, TypeError):
            net_weight_val = None

    kantar_weight_val = None
    if 'kantar_weight' in data and data['kantar_weight'] is not None and str(data['kantar_weight']).strip() != '':
        try:
            kantar_weight_val = float(data['kantar_weight'])
        except (ValueError, TypeError):
            kantar_weight_val = None

    result = database.complete_unloading(
        vehicle_id,
        unloaded_amount=unloaded_amount,
        unloaded_note=unloaded_note,
        dock_number=dock_number or None,
        net_weight=net_weight_val,
        kantar_weight=kantar_weight_val
    )
    return jsonify({'success': True, 'data': result, 'message': 'Boşaltma girişi kaydedildi.'})

@app.route('/api/public/approved-vehicles', methods=['GET'])
def api_public_approved_vehicles():
    """Şoförlerin takip edebileceği, liman giriş onayı ve boşaltma girişi yapılmış araçlar listesi (Herkese açık)"""
    limit = request.args.get('limit', default=50, type=int)
    tracking_data = database.get_driver_tracking_vehicles(limit=limit)
    return jsonify({
        'success': True,
        'port_vehicles': tracking_data['port_vehicles'],
        'unloaded_vehicles': tracking_data['unloaded_vehicles'],
        'vehicles': tracking_data['unloaded_vehicles']
    })

@app.route('/api/vehicles/<int:vehicle_id>', methods=['DELETE'])
def api_delete_vehicle(vehicle_id):
    """Araç kaydını silme"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401

    database.delete_vehicle(vehicle_id)
    return jsonify({'success': True, 'message': 'Kayıt silindi.'})

@app.route('/api/companies', methods=['GET'])
def api_get_companies():
    """Firma listesi"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401
    return jsonify({'success': True, 'companies': database.get_distinct_companies()})

@app.route('/api/stats', methods=['GET'])
def api_get_stats():
    """Operasyon paneli özet sayaçları"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401
    return jsonify({'success': True, 'stats': database.get_summary_stats()})

@app.route('/api/export/excel', methods=['GET'])
def api_export_excel():
    """Filtrelere göre Excel (.xlsx) raporu indirme"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401

    status = request.args.get('status')
    company = request.args.get('company')
    date_start = request.args.get('date_start')
    date_end = request.args.get('date_end')
    search = request.args.get('search')
    is_active_only = request.args.get('is_active', 'false').lower() == 'true'

    vehicles = database.get_vehicles(
        status=status,
        company=company,
        date_start=date_start,
        date_end=date_end,
        search=search,
        is_active_only=is_active_only
    )

    # Excel Oluşturma
    wb = openpyxl.Workbook()
    ws = wb.active
    is_archive_export = (status == 'BOSALTILDI')
    ws.title = "Boşaltılanlar Arşiv" if is_archive_export else "Tır Raporu"

    # Başlık stili
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    headers = [
        "Sıra", "Çekici Plaka", "Dorse Plaka", "Alıcı Firma", "Varış Tarihi/Saati", "Sisteme Kayıt",
        "Boşaltma Deposu", "Net Ağırlık (kg)", "Brüt Ağırlık (kg)", "Kantar Miktarı (kg)", "Kap Adedi", "Boşaltma Miktarı",
        "Boşaltma Giriş Zamanı", "Liman Giriş Onayı", "Liman Onay Zamanı", "Tescil Durumu", "Tescil Zamanı",
        "T1 / Tescil / Beyanname No", "Belge Fotoğrafı", "Mükerrer (3 Gün)", "Önemli Notlar", "Boşaltma Notu"
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    for idx, v in enumerate(vehicles, start=1):
        liman_str = "Onaylandı" if ('is_port_entered' in v and v['is_port_entered']) else "Bekliyor"
        liman_time = v.get('port_entered_at') or "-"
        tescil_str = "Yapıldı" if v['is_registered'] else "Bekliyor"
        bosaltma_str = "Giriş Yapıldı" if v['is_unloaded'] else "Bekliyor"
        belge_str = "Yüklendi (Mevcut)" if ('document_photo' in v and v['document_photo']) else "-"
        created_str = v.get('created_at') or "-"
        dup_str = "MÜKERRER (3 Gün)" if v.get('is_duplicate_3days') else "-"
        
        row_data = [
            idx,
            v['tractor_plate'],
            v['trailer_plate'],
            v['company_name'],
            v['arrival_date'],
            created_str,
            v['dock_number'] or "-",
            v.get('net_weight') if v.get('net_weight') is not None else "-",
            v['gross_weight'] if v['gross_weight'] is not None else "-",
            v.get('kantar_weight') if v.get('kantar_weight') is not None else "-",
            v['package_count'] if v['package_count'] is not None else "-",
            v['unloaded_amount'] or "-",
            v['unloaded_at'] or "-",
            liman_str,
            liman_time,
            tescil_str,
            v['registered_at'] or "-",
            v['customs_doc_no'] or "-",
            belge_str,
            dup_str,
            v['important_notes'] if ('important_notes' in v and v['important_notes']) else "-",
            v['unloaded_note'] or "-"
        ]
        ws.append(row_data)
        row_num = idx + 1
        ws.row_dimensions[row_num].height = 22

        # Renklendirme ve border
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [1, 2, 3, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

            # Boşaltıldıysa hafif yeşil satır
            if v['is_unloaded']:
                cell.fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")

    # Sütun genişliklerini otomatik ayarla
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    file_date = datetime.now().strftime('%Y%m%d_%H%M')
    if is_archive_export:
        filename = f"Tir_Bosaltilan_Arsiv_{file_date}.xlsx"
    else:
        filename = f"Tir_Bosaltma_Raporu_{file_date}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

# ==========================================
# VERİTABANI YÖNETİMİ & YEDEKLEME APILERI
# ==========================================

@app.route('/api/database/status', methods=['GET'])
def api_database_status():
    """Veritabanı ve yüklenen dosya depolama istatistiklerini getirir"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401
    try:
        stats = database.get_db_stats()
        return jsonify({'success': True, 'stats': stats})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@app.route('/api/database/backup/db', methods=['GET'])
def api_database_backup_db():
    """PostgreSQL kullanıldığı için eski SQLite .db yedekleme endpoint'i devre dışıdır."""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401

    return jsonify({
        'success': False,
        'message': 'Bu sistem PostgreSQL kullanmaktadır. Veritabanı yedeği için JSON yedekleme seçeneğini kullanınız.',
        'backup_endpoint': '/api/database/backup/json'
    }), 410

@app.route('/api/database/backup/json', methods=['GET'])
def api_database_backup_json():
    """Tüm araç ve ayar kayıtlarını JSON dosyası olarak indirir"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401
    try:
        import json
        data = database.export_all_data()
        json_str = json.dumps(data, ensure_ascii=False, indent=2)
        output = BytesIO(json_str.encode('utf-8'))
        now_str = datetime.now().strftime('%Y%m%d_%H%M')
        download_name = f"logistics_yedek_{now_str}.json"
        return send_file(
            output,
            as_attachment=True,
            download_name=download_name,
            mimetype="application/json"
        )
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@app.route('/api/database/cleanup', methods=['POST'])
def api_database_cleanup():
    """Veritabanını temizler veya sıfırlar (Güvenlik onay kodu ile)"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401
    
    data = request.get_json() or {}
    action = data.get('action') # 'archive' veya 'all'
    confirm_code = data.get('confirm_code', '').strip().upper()

    try:
        if action == 'archive':
            result = database.clear_archived_vehicles()
            return jsonify({
                'success': True,
                'message': f"Arşivdeki {result['deleted_vehicles']} araç ve {result['deleted_photos']} evrak fotoğrafı başarıyla temizlendi."
            })
        elif action == 'all':
            if confirm_code != 'SIFIRLA':
                return jsonify({
                    'success': False,
                    'message': "Tüm veritabanını sıfırlamak için onay kutusuna 'SIFIRLA' yazmanız gerekmektedir."
                }), 400
            
            result = database.reset_all_vehicles()
            return jsonify({
                'success': True,
                'message': f"Veritabanı tamamen sıfırlandı. ({result['deleted_vehicles']} araç silindi, fabrika ayarlarına dönüldü.)"
            })
        else:
            return jsonify({'success': False, 'message': 'Geçersiz temizleme işlemi.'}), 400
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@app.route('/api/database/restore', methods=['POST'])
def api_database_restore():
    """Yedek dosyasını (.db veya .json) geri yükler"""
    if not is_authenticated():
        return jsonify({'error': 'Unauthorized'}), 401

    if 'backup_file' not in request.files:
        return jsonify({'success': False, 'message': 'Lütfen geri yüklenecek yedek dosyasını seçiniz.'}), 400

    file = request.files['backup_file']
    if not file or not file.filename:
        return jsonify({'success': False, 'message': 'Geçerli bir dosya seçilmedi.'}), 400

    filename = file.filename.lower()
    try:
        if filename.endswith('.json'):
            import json
            content = file.read().decode('utf-8')
            json_data = json.loads(content)
            imported_count = database.import_data_from_json(json_data)
            return jsonify({
                'success': True,
                'message': f"JSON yedeği başarıyla geri yüklendi! ({imported_count} araç kaydı aktarıldı)"
            })
        elif filename.endswith('.db') or filename.endswith('.sqlite3'):
            # Güvenlik için SQLite başlığını doğrula
            header = file.read(16)
            if header != b"SQLite format 3\x00":
                return jsonify({'success': False, 'message': 'Yüklenen dosya geçerli bir SQLite veritabanı değil.'}), 400
            file.seek(0)

            # Mevcut veritabanının güvenlik yedeğini al
            safety_backup = database.DB_PATH + f".bak_{int(time.time())}"
            if os.path.exists(database.DB_PATH):
                import shutil
                shutil.copy2(database.DB_PATH, safety_backup)

            # Yeni veritabanı dosyasını kaydet
            file.save(database.DB_PATH)
            database.init_db()

            return jsonify({
                'success': True,
                'message': 'SQLite veritabanı başarıyla geri yüklendi ve aktifleştirildi.'
            })
        else:
            return jsonify({'success': False, 'message': 'Yalnızca .db veya .json uzantılı yedek dosyaları desteklenir.'}), 400
    except Exception as e:
        return jsonify({'success': False, 'message': f'Geri yükleme hatası: {str(e)}'}), 500

if __name__ == '__main__':
    local_ip = get_local_ip()
    port = int(os.environ.get('PORT', 8000))
    print("\n" + "=" * 60)
    print("🚀 TIR VE YÜK BOŞALTMA TAKİP SİSTEMİ BAŞLATILDI")
    print("=" * 60)
    print(f"💻 Bu Bilgisayardan Giriş:")
    print(f"   👉 Şoför Formu:     http://localhost:{port}/")
    print(f"   👉 Operasyon Paneli: http://localhost:{port}/operasyon")
    print(f"\n📱 Telefondan / Diğer Cihazlardan Giriş (Aynı Wi-Fi):")
    print(f"   👉 Şoför Formu:     http://{local_ip}:{port}/")
    print(f"   👉 Operasyon Paneli: http://{local_ip}:{port}/operasyon")
    print(f"\n🔑 Operasyon Kullanıcı Adı: {OPERATIONS_USERNAME}")
    print("=" * 60 + "\n")

    app.run(host='0.0.0.0', port=port, debug=False)
